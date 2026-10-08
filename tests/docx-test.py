#!/usr/bin/env python3
"""docx 腳本的自動化測試：在暫存目錄合成一份帶自訂樣式的範本，驗證轉換與傾印。
執行：python -X utf8 tests/docx-test.py"""
import json
import os
import subprocess
import sys
import tempfile

import docx
from docx.enum.style import WD_STYLE_TYPE

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, "..", "scripts", "docx")
fails = []


def check(name, cond, detail=""):
    print(("PASS  " if cond else "FAIL  ") + name + ("" if cond else "  " + str(detail)))
    if not cond:
        fails.append(name)


def run(script, *args):
    return subprocess.run([sys.executable, "-X", "utf8", os.path.join(SCRIPTS, script), *args],
                          capture_output=True, text=True, encoding="utf-8")


with tempfile.TemporaryDirectory() as tmp:
    tpl = os.path.join(tmp, "tpl.docx")
    d = docx.Document()
    for name in ("樣式H1", "樣式H2", "樣式H3", "樣式H4", "樣式內文"):
        d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("範本殘留的範例文字", style="樣式內文")
    d.sections[0].header.paragraphs[0].text = "範本頁首"
    d.save(tpl)

    mapping = os.path.join(tmp, "map.json")
    with open(mapping, "w", encoding="utf-8") as f:
        json.dump({"template": "tpl.docx",
                   "styles": {"h1": "樣式H1", "h2": "樣式H2", "h3": "樣式H3", "h4": "樣式H4", "body": "樣式內文"},
                   "skip_sections": ["Release Info"], "drop_line_patterns": [r"^\*\*對應需求\*\*"]}, f)
    md = os.path.join(tmp, "a.md")
    with open(md, "w", encoding="utf-8") as f:
        f.write("# 功能說明\n## 功能概述\n以文字概述 **重點**。\n\n- 項目一\n  - 子項\n\n"
                "| 欄位 | 說明 |\n|---|---|\n| A | 甲 |\n\n```\nSELECT 1\n```\n\n"
                "**對應需求**：Req 1.1\n![畫面](不存在.png)\n\n## Release Info\n- **Status**: draft\n\n# 功能規格\n內文結尾\n")
    out = os.path.join(tmp, "o.docx")

    r = run("md_to_docx.py", "--mapping", mapping, "--md", md, "--out", out)
    check("轉換成功", r.returncode == 0, r.stderr)
    dump = run("docx_dump.py", out).stdout
    check("標題套用對應樣式", "[樣式H1] 功能說明" in dump and "[樣式H2] 功能概述" in dump, dump)
    check("範本殘留內文已清除", "範本殘留" not in dump, dump)
    check("表格輸出", "<TABLE>" in dump and "A | 甲" in dump, dump)
    check("程式碼區塊保留", "SELECT 1" in dump, dump)
    check("圖片缺檔以預留文字取代", "請貼上圖片：畫面" in dump, dump)
    check("skip_sections 略過整段", "Release Info" not in dump and "Status" not in dump, dump)
    check("略過後的下一個同級標題恢復輸出", "[樣式H1] 功能規格" in dump, dump)
    check("drop_line_patterns 濾掉內部追溯行", "對應需求" not in dump and "Req 1.1" not in dump, dump)
    check("範本頁首保留", docx.Document(out).sections[0].header.paragraphs[0].text == "範本頁首")

    r = run("md_to_docx.py", "--mapping", mapping, "--md", md, "--out", out)
    check("輸出已存在時拒絕覆蓋", r.returncode != 0 and "不覆蓋" in r.stderr, r.stderr)

    with open(mapping, "w", encoding="utf-8") as f:
        json.dump({"template": "tpl.docx",
                   "styles": {"h1": "不存在的樣式", "h2": "樣式H2", "h3": "樣式H3", "h4": "樣式H4", "body": "樣式內文"}}, f)
    r = run("md_to_docx.py", "--mapping", mapping, "--md", md, "--out", os.path.join(tmp, "o2.docx"))
    check("樣式不存在時報錯並列出現有樣式", r.returncode != 0 and "不存在的樣式" in r.stderr and "樣式H2" in r.stderr, r.stderr)

    # ---- 品保回報的邊界案例 ----
    with open(mapping, "w", encoding="utf-8") as f:
        json.dump({"template": "tpl.docx",
                   "styles": {"h1": "樣式H1", "h2": "樣式H2", "h3": "樣式H3", "h4": "樣式H4", "body": "樣式內文"},
                   "skip_sections": ["內部"]}, f)
    md2 = os.path.join(tmp, "b.md")
    with open(md2, "w", encoding="utf-8-sig") as f:
        f.write("# 標題一\n```bash\n# 這是 shell 註解不是標題\necho 1\n```\n"
                "## 內部\n```\n# 內部區塊裡的註解\n```\n洩漏測試句\n## 後續\n<!-- 內部備註 -->\n正常內文\n"
                "|-|-|\n| 欄 | 值 \\| 含管線 |\n|-|-|\n| a | b |\n")
    out2 = os.path.join(tmp, "b.docx")
    r = run("md_to_docx.py", "--mapping", mapping, "--md", md2, "--out", out2)
    check("邊界案例轉換成功（含 BOM）", r.returncode == 0, r.stderr)
    dump2 = run("docx_dump.py", out2).stdout
    check("程式碼區塊內的 # 行不被當標題", "[樣式H1] 這是 shell 註解" not in dump2 and "# 這是 shell 註解不是標題" in dump2, dump2)
    check("BOM 不讓第一個標題變內文", "[樣式H1] 標題一" in dump2, dump2)
    check("略過區段內含程式碼區塊也整段被擋", "內部區塊裡的註解" not in dump2 and "洩漏測試句" not in dump2, dump2)
    check("略過結束後恢復輸出", "[樣式H2] 後續" in dump2, dump2)
    check("HTML 註解不帶進輸出", "內部備註" not in dump2, dump2)
    check("跳脫的管線留在儲存格內", "值 | 含管線" in dump2, dump2)

    sub = os.path.join(tmp, "sub")
    os.makedirs(sub)
    r = run("md_to_docx.py", "--mapping", mapping, "--md", md2, "--out", os.path.join(tmp, "nodir", "x.docx"))
    check("輸出目錄不存在給可行動的訊息", r.returncode != 0 and "輸出目錄不存在" in r.stderr and "Traceback" not in r.stderr, r.stderr)
    r = run("md_to_docx.py", "--mapping", mapping, "--md", os.path.join(tmp, "不存在.md"), "--out", os.path.join(tmp, "y.docx"))
    check("來源 md 不存在給可行動的訊息", r.returncode != 0 and "找不到來源" in r.stderr and "Traceback" not in r.stderr, r.stderr)
    bad = os.path.join(tmp, "bad.json")
    with open(bad, "w", encoding="utf-8") as f:
        f.write("{壞掉")
    r = run("md_to_docx.py", "--mapping", bad, "--md", md2, "--out", os.path.join(tmp, "z.docx"))
    check("對應檔 JSON 壞掉給可行動的訊息", r.returncode != 0 and "JSON" in r.stderr and "Traceback" not in r.stderr, r.stderr)
    empty = os.path.join(tmp, "empty.md")
    open(empty, "w").close()
    r = run("md_to_docx.py", "--mapping", mapping, "--md", empty, "--out", os.path.join(tmp, "e.docx"))
    check("空 md 拒絕產出", r.returncode != 0 and "空的" in r.stderr, r.stderr)
    r = run("docx_dump.py", os.path.join(tmp, "不存在.docx"))
    check("docx_dump 對不存在的檔給可行動的訊息", r.returncode != 0 and "Traceback" not in r.stderr and "找不到" in r.stderr, r.stderr)

    # 圖片相對路徑以 md 所在目錄為準，與執行時的 cwd 無關
    import base64
    with open(os.path.join(sub, "p.png"), "wb") as f:
        f.write(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))
    md3 = os.path.join(sub, "c.md")
    with open(md3, "w", encoding="utf-8") as f:
        f.write("# 圖\n![小圖](p.png)\n")
    out3 = os.path.join(tmp, "c.docx")
    r = subprocess.run([sys.executable, "-X", "utf8", os.path.join(SCRIPTS, "md_to_docx.py"), "--mapping", mapping,
                        "--md", md3, "--out", out3], capture_output=True, text=True, encoding="utf-8", cwd=tmp)
    check("圖片路徑相對於 md 目錄", r.returncode == 0 and len(docx.Document(out3).inline_shapes) == 1, r.stderr)


print("\n全部通過" if not fails else "\n失敗：%d 項" % len(fails))
sys.exit(1 if fails else 0)
